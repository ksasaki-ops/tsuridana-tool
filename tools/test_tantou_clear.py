"""
担当者 (選択の記憶・Excel 記載) と入力クリア機能のテスト

使い方:
    py -3 tools/test_tantou_clear.py
"""
import os
import sys
import tempfile
import tkinter as tk
from datetime import date

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.argv[0] = os.path.join(ROOT, "tsuridana.py")

from openpyxl import load_workbook

from core.excel_writer import generate_excel
from core.text_builder import OrderSpec
from gui.main_window import MainWindow
from utils.settings import load_settings, remember_tantou, save_settings


def _spec(**kw):
    base = dict(
        customer="テスト", property_name="テスト物件",
        W=760, D=450, H=600,
        men_zai_ari=False, tobira_hinban="",
        body_material="白ポリ", filler="両側",
        hg=False, tobira_enchou=False, kirikake=False,
    )
    base.update(kw)
    return OrderSpec(**base)


tmp = tempfile.mkdtemp(prefix="tsuridana_test_")
settings_path = os.path.join(tmp, "sub", "settings.json")

# === テスト1: 設定ファイルが無い / 壊れている場合は初期値 ===
print("=== テスト1: 設定の初期値とフォールバック ===")
s = load_settings(settings_path)
assert s == {"tantou": "", "tantou_list": []}, f"FAIL: 初期値: {s}"
os.makedirs(os.path.dirname(settings_path))
with open(settings_path, "w", encoding="utf-8") as fp:
    fp.write("{ broken json")
assert load_settings(settings_path) == {"tantou": "", "tantou_list": []}, "FAIL: 壊れた JSON"
with open(settings_path, "w", encoding="utf-8") as fp:
    fp.write('{"tantou": 123, "tantou_list": "x"}')
assert load_settings(settings_path) == {"tantou": "", "tantou_list": []}, "FAIL: 型が不正"
print("  OK: 無い・壊れている・型不正 → 初期値")

# === テスト2: 担当者の記憶 (重複なし・最新が先頭・空は無視) ===
print("=== テスト2: remember_tantou / save / load ===")
s = load_settings(settings_path)
s = remember_tantou(s, "佐々木")
s = remember_tantou(s, " 中三川 ")
s = remember_tantou(s, "佐々木")
s = remember_tantou(s, "   ")
assert s == {"tantou": "佐々木", "tantou_list": ["佐々木", "中三川"]}, f"FAIL: {s}"
save_settings(s, settings_path)
assert load_settings(settings_path) == s, "FAIL: 保存→読込で一致しない"
print("  OK: 重複なし・最新先頭・空無視・ラウンドトリップ")

# === テスト3: Excel の K9 に「担当：○○」 ===
print("=== テスト3: Excel 記載 ===")
out = generate_excel(_spec(tantou="佐々木"), os.path.join(tmp, "with.xlsx"))
ws = load_workbook(out).active
assert ws["K9"].value == "担当：佐々木", f"FAIL: K9={ws['K9'].value!r}"
out = generate_excel(_spec(), os.path.join(tmp, "without.xlsx"))
ws = load_workbook(out).active
assert ws["K9"].value is None, f"FAIL: 担当者なしで K9={ws['K9'].value!r}"
print("  OK: 担当者あり→K9 記載 / なし→空のまま")

# === テスト4: GUI — 次回起動で担当者が復元される ===
print("=== テスト4: GUI 担当者の復元 ===")
root = tk.Tk()
root.withdraw()
win = MainWindow(root, settings_path=settings_path)
assert win.var_tantou.get() == "佐々木", f"FAIL: 復元: {win.var_tantou.get()!r}"
assert list(win.combo_tantou["values"]) == ["佐々木", "中三川"], f"FAIL: 候補: {win.combo_tantou['values']}"
win.var_tantou.set("杉原")
win._remember_tantou()
root.destroy()

root = tk.Tk()
root.withdraw()
win = MainWindow(root, settings_path=settings_path)
assert win.var_tantou.get() == "杉原", f"FAIL: 再起動後: {win.var_tantou.get()!r}"
assert list(win.combo_tantou["values"]) == ["杉原", "佐々木", "中三川"], f"FAIL: {win.combo_tantou['values']}"
print("  OK: 選択・手入力した担当者が次回起動で復元")

# === テスト5: GUI — クリアで初期状態に戻り、担当者だけ残る ===
print("=== テスト5: 入力クリア ===")
win.var_customer.set("山田")
win.var_property.set("テストマンション")
win.var_hassou_no.set("A-001")
win.var_hassou_date.set("2020-01-01")
win.var_W.set("760"); win.var_D.set("450"); win.var_H.set("600")
win.var_men_zai.set(False); win.var_tobira_hinban.set("XX-1")
win.var_material.set("グレーオーク")
win.var_tana_count.set("3")
win.var_filler.set("なし")
win.var_hg.set(False)
win.var_tobira_enc.set(True)
win.var_kirikake.set(True); win.var_kk_W.set("100"); win.var_kk_H.set("50")
win.var_bikou.set("メモ")
win._toggle_men_zai(); win._toggle_kirikake()

win._reset_fields()

expected = {
    "var_customer": "", "var_property": "", "var_hassou_no": "",
    "var_hassou_date": date.today().strftime("%Y-%m-%d"),
    "var_W": "", "var_D": "", "var_H": "",
    "var_men_zai": True, "var_tobira_hinban": "",
    "var_material": "白ポリ", "var_tana_count": "1", "var_filler": "両側",
    "var_hg": True, "var_tobira_enc": False,
    "var_kirikake": False, "var_kk_W": "", "var_kk_H": "",
    "var_bikou": "",
}
for name, want in expected.items():
    got = getattr(win, name).get()
    assert got == want, f"FAIL: {name}: {got!r} != {want!r}"
assert win.var_tantou.get() == "杉原", "FAIL: クリアで担当者が消えた"
assert str(win.entry_tobira_hinban["state"]) == "normal", "FAIL: 扉品番欄の状態"
assert str(win.entry_kk_W["state"]) == "disabled", "FAIL: 切り欠き欄の状態"
root.destroy()
print("  OK: 全項目が初期状態・発注日は当日・担当者は保持・欄の有効/無効も初期化")

print("\n全テスト PASS")
